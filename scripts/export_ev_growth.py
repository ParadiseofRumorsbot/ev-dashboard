"""Export recalculated EV workbook values to the dashboard. No forecasting formulas here.

Run after recalculating and saving the workbook in Excel:
    python scripts/export_ev_growth.py --workbook PATH
The dashboard is a published snapshot, not a live Excel connection.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument('--workbook', required=True, type=Path)
parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'data/ev_growth.json')
args = parser.parse_args()
w = load_workbook(args.workbook, data_only=True)
formulas = load_workbook(args.workbook, data_only=False)
for sheet in formulas:
    for row in sheet:
        for c in row:
            value = w[sheet.title][c.coordinate]
            if c.data_type == 'f' and (value.value is None or value.data_type == 'e'):
                raise ValueError(f'Recalculate and save workbook first: {sheet.title}!{c.coordinate}')

def rows(name, start, end, cols):
    return [list(r) for r in w[name].iter_rows(min_row=start, max_row=end, max_col=cols, values_only=True)]

def source_url(v):
    urls=re.findall(r'https://[^\s;]+', v or '')
    return urls[0] if urls else None

suppliers = {'Samsung SDI': '삼성SDI', 'LG Energy Solution': 'LG에너지솔루션', 'SK On': 'SK온'}
reverse = {v:k for k,v in suppliers.items()}
ms, cs, ps = w['차종별 계산'], w['셀사별 배정'], w['출시·수주 확인']
assert ms['AA6'].value == '2026E 산출 대', 'Unsupported workbook schema'
assert cs['X6'].value == '2026E 셀 GWh', 'Unsupported workbook schema'
assert w['공급변경 시나리오']['B4'].value == '기준안', 'Save the workbook with 기준안 selected before publishing'
data = {
    'version': '2026-09-30-growth-v1', 'asOf': '2026-09-30', 'actualThrough': '2026-06',
    'workbookSha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
    'metric': '판매차량 탑재 배터리 수요(GWh)',
    'baseline': '2026E 상반기 단순 연환산. Tesla 3/Y는 2026-06-26 회사 취합 컨센서스 우선.',
    'future': '2028·2029: 기존 최근12개월 유지 또는 공개 전망. 신규 물량 미확인 제외.',
    'years': [2026,2028,2029], 'suppliers': suppliers,
    'models': [], 'cells': [], 'contracts': [], 'capacity': [], 'factories': [], 'utilization': [], 'sources': [],
    'summary': rows('셀3사 요약',8,16,10),
    'caseRates': rows('공급변경 시나리오',10,13,3),
}
for i,r in enumerate(rows(ms.title,7,ms.max_row,30),7):
    if not r[0]: continue
    data['models'].append(dict(id=r[0],region=r[1],country=r[2],oem=r[3],model=r[4],power=r[5],platform=r[6],
        units=[r[26],r[17],r[20]],gwh=[r[27],r[18],r[21]],ttm=r[14],h1=r[13],method=r[28],action28=r[16],
        action29=r[19],status=r[22],evidence=r[23],reason=r[24],sheetRow=i))
for i,r in enumerate(rows(cs.title,7,cs.max_row,29),7):
    if not r[0]: continue
    data['cells'].append(dict(modelId=r[0],region=r[1],supplier=r[5],gwh=[r[23],r[13],r[15]],ttm=r[9],h1=r[8],hold=[r[16],r[17]],sheetRow=i))
data['transfers']=[dict(modelId=r[0],region=r[1],model=r[3],**{'from':r[4],'to':r[5]},gwh=[r[6],r[7]])
    for r in rows('공급변경 시나리오',20,formulas['공급변경 시나리오'].max_row,12) if r[0]]
for i,r in enumerate(rows(ps.title,7,ps.max_row,19),7):
    if not r[0]: continue
    data['contracts'].append(dict(region=r[0],oem=r[1],model=r[2],supplier=r[3],status=r[4],gwh=[r[5],r[6]],contractGwh=r[7],
        period=r[8],start=r[9],power=r[10],platform=r[11],pack=r[12],note=r[13],evidence=r[15],reference=r[17],referenceUnit=r[18],sheetRow=i))
for r in rows('CAPA·가동률',8,13,9):
    data['capacity'].append(dict(supplier=reverse[r[0]],region=r[1],capacity=r[2:5],effectiveEV2026=r[5],matchedDemand2026=r[6],demandRatio2026=r[7],evUtilization2026=r[8],source='battery_scenario.html#cell3Card'))
for r in rows('CAPA·가동률',18,24,11):
    data['utilization'].append(dict(supplier=reverse[r[0]],region=r[1],plant=r[2],product=r[3],period=r[4],value=r[5],kind=r[6],
        apply=r[7],denominator=r[8],note='사용자 전달 추정. 원문·집계 범위·작성일 미확인. 계산 미적용.' if r[10]=='U04' else r[9],source=r[10]))
for r in rows('CAPA·가동률',30,w['CAPA·가동률'].max_row,15):
    if r[0] not in reverse: continue
    data['factories'].append(dict(supplier=reverse[r[0]],region=r[1],name=r[2],app=r[3],capacity=r[4:7],status=r[7],note=r[10],client=r[11]))
for r in rows('공개 전망·근거',19,62,10):
    data['sources'].append(dict(id=r[0],name=r[1],date=r[2],kind=r[3],fact=r[5],use=r[6],url=source_url(r[9]),urls=re.findall(r'https://[^\s;]+',r[9] or '')))
for r in rows('CAPA·가동률',30,w['CAPA·가동률'].max_row,5):
    if r[0] not in ['U01','U02','U03','U04','U05']: continue
    data['sources'].append(dict(id=r[0],name='사용자 전달 추정' if r[0]=='U04' else r[1],date=r[2],kind=r[3],url=source_url(r[4])))
args.output.parent.mkdir(parents=True,exist_ok=True)
args.output.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print(json.dumps(dict(models=len(data['models']),cells=len(data['cells']),contracts=len(data['contracts']),factories=len(data['factories']),output=str(args.output)),ensure_ascii=False))
